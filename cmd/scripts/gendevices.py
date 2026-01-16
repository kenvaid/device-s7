#!/usr/bin/env python3
"""
Generate a devices YAML file from a template YAML and an Excel file.

Usage:
  python3 gendevices.py -t template.yml -e devices.xlsx -o out.yml

Behavior:
  - Reads the template YAML and uses the first item under `deviceList` as a template.
  - Reads the first sheet (or a named sheet) from the Excel workbook.
  - For each non-empty row (starting at `--start-row`, default 1), reads:
      Column A -> device `name`
      Column B -> protocol `Port`
    Other fields are copied from the template without change.

Dependencies: pyyaml, openpyxl
  pip install pyyaml openpyxl
"""
import argparse
import copy
import re
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook


def load_template(path: Path):
    with path.open('r', encoding='utf-8') as f:
        data = yaml.safe_load(f)
    if not isinstance(data, dict) or 'deviceList' not in data or not data['deviceList']:
        raise ValueError('Template YAML must contain a non-empty `deviceList`')
    return data


def update_protocol_port(protocols: dict, port_value):
    # Try to update the Port under a known protocol key (e.g. modbus-tcp).
    for proto_name, proto_val in protocols.items():
        if isinstance(proto_val, dict) and 'Port' in proto_val:
            proto_val['Port'] = int(port_value) if port_value is not None and str(port_value).strip() != '' else proto_val.get('Port')
            return True
    # If no explicit 'Port' key found, try to set Port under the first protocol if it's a dict
    for proto_name, proto_val in protocols.items():
        if isinstance(proto_val, dict):
            proto_val['Port'] = int(port_value) if port_value is not None and str(port_value).strip() != '' else proto_val.get('Port')
            return True
    return False


def preprocess_name(raw_name, suffix: str | None):
    """Remove '-', '_', '.'; lowercase; append suffix with '-' if provided."""
    s = '' if raw_name is None else str(raw_name)
    # remove specified characters
    s = re.sub(r'[-_.]', '', s)
    s = s.lower()
    if suffix is None or str(suffix).strip() == '':
        return s
    return f"{s}-{str(suffix).strip()}"


def generate_devices(template_data: dict, excel_path: Path, sheet_name: str = None, start_row: int = 1, suffix: str | None = None):
    wb = load_workbook(filename=str(excel_path), data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]

    template_device = template_data['deviceList'][0]
    output_devices = []

    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if row is None:
            continue
        # Column A -> name, Column B -> Port
        name = row[0] if len(row) >= 1 else None
        port = row[1] if len(row) >= 2 else None


        if name is None or (isinstance(name, str) and name.strip() == ''):
            # skip empty name rows
            continue

        device = copy.deepcopy(template_device)
        device['name'] = preprocess_name(name, suffix)

        # Update Port inside protocols (if present)
        if 'protocols' in device and isinstance(device['protocols'], dict):
            updated = update_protocol_port(device['protocols'], port)
            if not updated and port is not None:
                # As a fallback, add a modbus-tcp entry
                device.setdefault('protocols', {})['modbus-tcp'] = {'Port': int(port)}

        output_devices.append(device)

    return {'deviceList': output_devices}


def write_output(data: dict, out_path: Path):
    # Dump YAML with preserved mapping order and readable style
    with out_path.open('w', encoding='utf-8') as f:
        yaml.safe_dump(data, f, sort_keys=False, allow_unicode=True)


def main():
    parser = argparse.ArgumentParser(description='Generate devices YAML from template and Excel')
    parser.add_argument('-t', '--template', required=True, help='Path to template YAML (with deviceList)')
    parser.add_argument('-e', '--excel', required=True, help='Path to Excel file (.xlsx)')
    parser.add_argument('-o', '--output', required=True, help='Path to output YAML file')
    parser.add_argument('--sheet', help='Excel sheet name (defaults to first sheet)')
    parser.add_argument('--start-row', type=int, default=1, help='1-based first row to read (default: 1)')
    parser.add_argument('-s', '--suffix', help='Suffix to append to name (joined with "-")', default='')

    args = parser.parse_args()

    template_path = Path(args.template)
    excel_path = Path(args.excel)
    out_path = Path(args.output)

    if not template_path.exists():
        print(f'Template file not found: {template_path}', file=sys.stderr)
        sys.exit(2)
    if not excel_path.exists():
        print(f'Excel file not found: {excel_path}', file=sys.stderr)
        sys.exit(2)

    try:
        template = load_template(template_path)
        result = generate_devices(template, excel_path, sheet_name=args.sheet, start_row=args.start_row, suffix=args.suffix)
        write_output(result, out_path)
        print(f'Wrote {len(result.get("deviceList", []))} devices to {out_path}')
    except Exception as e:
        print('Error:', str(e), file=sys.stderr)
        sys.exit(1)


if __name__ == '__main__':
    main()
